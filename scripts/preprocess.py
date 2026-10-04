"""Explicit source mappings; evidence registration and conservative candidates."""
from __future__ import annotations

import argparse
import csv
import shutil
from collections import Counter
from datetime import datetime
from pathlib import Path

from common import (VERSION, check_contract, envelope, money, period_dir, period_key,
                    period_of, read_json, relative_file, sha, stable_id, write_json)


def scalar(value):
    return value.isoformat() if isinstance(value, datetime) else value


def source_rows(path, mapping):
    fmt = mapping['format']
    if fmt == 'standardized_json':
        value = read_json(path)
        rows = value['records']
        if not isinstance(rows, list) or any(not isinstance(r, dict) for r in rows):
            raise ValueError('统一 JSON 接入须为 {records: [字段映射行]}')
        return [(None, i + 1, r) for i, r in enumerate(rows)]
    if fmt in ('csv', 'tsv'):
        with Path(path).open(encoding=mapping.get('encoding', 'utf-8-sig'), newline='') as stream:
            rows = list(csv.reader(stream, delimiter=mapping.get('delimiter', '\t' if fmt == 'tsv' else ',')))
        sheets = [(None, rows)]
    elif fmt == 'xlsx':
        from openpyxl import load_workbook
        book = load_workbook(path, read_only=True, data_only=False)
        try:
            sheets = [(s.title, [list(row) for row in s.iter_rows(values_only=True)]) for s in book
                      if mapping.get('sheet') is None or s.title == mapping['sheet']]
        finally:
            book.close()
        if not sheets:
            raise ValueError('指定工作表不存在')
    else:
        raise ValueError('不可读取的格式')
    result = []
    header = mapping.get('header_row', 1)
    for sheet, rows in sheets:
        if not 1 <= header <= len(rows):
            raise ValueError('表头行越界')
        headers = [str(h).strip() if h is not None else '' for h in rows[header - 1]]
        if len(set(headers)) != len(headers) or any(not h for h in headers):
            raise ValueError('表头为空或重复，请由读取工具整理后接入统一 JSON')
        for i, row in enumerate(rows[header:], header + 1):
            if not any(v is not None and str(v).strip() for v in row):
                continue
            result.append((sheet, i, {h: scalar(row[j]) if j < len(row) else None for j, h in enumerate(headers)}))
    return result


def trade_time(value, formats):
    if isinstance(value, datetime):
        return value.isoformat(), 'datetime'
    value = str(value).strip()
    try:
        dt = datetime.fromisoformat(value)
    except ValueError:
        dt = None
        for fmt in formats:
            try:
                dt = datetime.strptime(value, fmt)
                break
            except ValueError:
                pass
        if dt is None:
            raise ValueError('交易日期无法按映射解析')
    return dt.isoformat(), 'datetime' if ':' in value else 'date'


def mapped_record(fields, mapping, location, currency, precision):
    columns = mapping['fields']
    allowed = {'trade_time','raw_amount','flow_direction','account_key','original_trade_no',
               'currency','raw_counterparty','raw_description','raw_status'}
    if columns.keys() - allowed or not {'trade_time','raw_amount','flow_direction'} <= columns.keys():
        raise ValueError('映射含未知字段或缺少日期、金额、方向')
    if any(column not in fields for column in columns.values()):
        raise ValueError('映射列在当前文件中不存在')
    values = {key: fields[column] for key, column in columns.items()}
    cur = str(values.get('currency') or mapping.get('currency') or currency)
    if cur != currency:
        raise ValueError('异币种记录，须分开分析')
    raw_amount = money(values['raw_amount'], precision)
    multiplier = mapping.get('amount_multiplier', 1)
    if multiplier <= 0:
        raise ValueError('金额换算倍率须为正，仅用于单位换算，不能自动换汇')
    from decimal import Decimal
    amount = money(raw_amount * Decimal(str(multiplier)), precision)
    if Decimal(str(float(amount))) != amount:
        raise ValueError('金额无法无损保存为 JSON 数值，请缩小分析单位或使用扩展精确数值契约')
    dt, time_precision = trade_time(values['trade_time'], mapping.get('date_formats', []))
    direction = mapping['direction_values'].get(str(values['flow_direction']).strip())
    if direction is None:
        raise ValueError('方向不在显式映射中')
    account = values.get('account_key') or mapping.get('account_key') or None
    txid = values.get('original_trade_no') or None
    rid = stable_id('RAW', mapping['original_source'], account, location['file'], location['sha256'], location['sheet'], location['row'])
    record = dict(record_id=rid,original_source=mapping['original_source'],account_key=str(account) if account else None,
                  original_trade_no=str(txid) if txid else None,trade_time=dt,time_precision=time_precision,
                  raw_amount=float(amount),flow_direction=direction,currency=cur,
                  raw_counterparty=str(values.get('raw_counterparty') or ''),
                  raw_description=str(values.get('raw_description') or ''),raw_status=str(values.get('raw_status') or ''),
                  raw_fields=fields,source_location=location)
    check_contract(record, 'raw_record')
    return record


def register(root, period, path):
    path = Path(path).resolve()
    content_hash = sha(path)
    target = period_dir(root, period) / '原始数据' / path.name
    target.parent.mkdir(parents=True, exist_ok=True)
    if target.exists() and sha(target) != content_hash:
        target = target.with_name(target.stem + '-' + content_hash[:12] + target.suffix)
    if target.resolve() != path and not target.exists():
        shutil.copyfile(path, target)
    relative_file(root, target.relative_to(Path(root).resolve()).as_posix())
    return target


def mapping_for(mappings, filename):
    return mappings.get('files', {}).get(filename) if 'files' in mappings else mappings


def intake(root, period=None, files=None, output=None, mapping=None, currency='CNY', amount_precision=None):
    root = Path(root).resolve()
    root.mkdir(parents=True, exist_ok=True)
    precision = 2 if currency == 'CNY' and amount_precision is None else amount_precision
    if precision is None or not isinstance(precision, int) or not 0 <= precision <= 8:
        raise ValueError('非默认币种须显式指定 0—8 位金额精度')
    if mapping is None:
        raise ValueError('必须提供当前文件的字段映射')
    files = list(dict.fromkeys(str(Path(p).resolve()) for p in files or []))
    if not period:
        if files:
            periods = set()
            for file in files:
                m = mapping_for(mapping, Path(file).name)
                if m:
                    check_contract(m,'source_mapping')
                    for _, _, row in source_rows(file, m):
                        try:
                            dt, _ = trade_time(row[m['fields']['trade_time']], m.get('date_formats', []))
                            periods.add(period_of(dt))
                        except (ValueError, KeyError):
                            pass
            if len(periods) != 1:
                raise ValueError('文件范围未能唯一确定账期，请显式指定 --period')
            period = periods.pop()
        else:
            eligible = []
            for d in root.iterdir():
                try:
                    period_key(d.name)
                except ValueError:
                    continue
                sources = list((d / '原始数据').glob('*'))
                prior = read_json(d / '处理结果/原始标准化.json', {})
                known = {m['file']: m['sha256'] for m in prior.get('source_manifest', [])}
                if sources and (not prior or any(known.get(p.relative_to(root).as_posix()) != sha(p) for p in sources if p.is_file())):
                    eligible.append(d.name)
            if not eligible:
                raise ValueError('没有未处理或源文件已变化的账期')
            period = max(eligible, key=period_key)
    dest = period_dir(root, period)
    if not files:
        files = [str(p) for p in sorted((dest / '原始数据').glob('*')) if p.is_file()]
    raw, rejected, manifests, needs_reading, normalized, conflicts = [], [], [], [], {}, []
    for filename in files:
        path = register(root, period, filename)
        rel = path.relative_to(root).as_posix()
        mapping_item = mapping_for(mapping, Path(filename).name)
        m = dict(file=rel,sha256=sha(path),original_name=Path(filename).name,
                 original_source=mapping_item.get('original_source') if mapping_item else None,
                 account_key=mapping_item.get('account_key') if mapping_item else None,role='structured' if path.suffix.lower()=='.json' else 'raw',
                 read_status='read',record_count=0,mapping=mapping_item)
        if mapping_item is None:
            m['read_status']='needs_reading'
            needs_reading.append(rel)
            manifests.append(m)
            continue
        check_contract(mapping_item,'source_mapping')
        try:
            rows = source_rows(path, mapping_item)
        except (ValueError, OSError, KeyError, csv.Error) as exc:
            m['read_status']='needs_reading'
            needs_reading.append(rel)
            manifests.append(m)
            conflicts.append(dict(file=rel,reason=str(exc)))
            continue
        for sheet, row, fields in rows:
            location = dict(file=rel,sha256=m['sha256'],sheet=sheet,row=row)
            try:
                r = mapped_record(fields,mapping_item,location,currency,precision)
            except ValueError as exc:
                rejected.append(dict(file=rel,sheet=sheet,row=row,reason=str(exc),raw_fields=fields))
                continue
            raw.append(r)
            m['record_count'] += 1
            nid = stable_id('REC',r['original_source'],r['account_key'],r['original_trade_no']) if (
                mapping_item['transaction_id_is_stable'] and r['original_trade_no'] and r['account_key']) else stable_id('REC',r['record_id'])
            n = {k:v for k,v in r.items() if k not in ('raw_fields','source_location')}
            n.update(record_id=nid,raw_record_ids=[r['record_id']],source_locations=[location])
            if nid in normalized:
                prior = normalized[nid]
                keys = set(n) - {'raw_record_ids','source_locations'}
                if any(n[k]!=prior[k] for k in keys):
                    nid=stable_id('REC',nid,r['record_id'])
                    n['record_id']=nid
                    conflicts.append(dict(record_ids=[prior['record_id'],nid],reason='稳定交易号对应事实冲突，保留两条供核查'))
                else:
                    prior['raw_record_ids'].append(r['record_id'])
                    prior['source_locations'].append(location)
                    continue
            normalized[nid] = n
        manifests.append(m)
    raw = sorted({r['record_id']:r for r in raw}.values(),key=lambda r:r['record_id'])
    selected = sorted(normalized.values(),key=lambda r:r['record_id'])
    candidates = []
    for i, a in enumerate(selected):
        for b in selected[i+1:]:
            if a['raw_amount']!=b['raw_amount'] or a['currency']!=b['currency']:
                continue
            delta=abs((datetime.fromisoformat(a['trade_time']).replace(tzinfo=None)-datetime.fromisoformat(b['trade_time']).replace(tzinfo=None)).total_seconds())
            if delta <= 86400:
                candidates.append(dict(record_ids=[a['record_id'],b['record_id']],match_type='重复候选' if a['flow_direction']==b['flow_direction'] else '退款/回款候选',status='candidate'))
    degree=Counter(rid for c in candidates for rid in c['record_ids'])
    for candidate in candidates:
        candidate['ambiguous']=any(degree[r]>1 for r in candidate['record_ids'])
    standardized = envelope(period,currency=currency,amount_precision=precision,raw_records=raw,
                            normalized_records=selected,source_manifest=manifests,rejected_rows=rejected)
    check_contract(standardized,'standardized')
    out = Path(output).resolve() if output else dest
    if not out.is_relative_to(root):
        raise ValueError('--output 必须在显式工作区内')
    pre=envelope(period,publication_status='draft',candidates=candidates,conflicts=conflicts,
                 needs_reading=needs_reading,rejected_rows=rejected,
                 outside_period_record_ids=[r['record_id'] for r in selected if period_of(r['trade_time'])!=period],
                 question_groups=[dict(topic='coverage',status='needs_context',record_ids=[]),
                                  dict(topic='economic_semantics',status='needs_review',record_ids=[r['record_id'] for r in selected if period_of(r['trade_time'])==period])])
    write_json(out/'处理结果/原始标准化.json',standardized)
    write_json(out/'处理结果/预分析.json',pre)
    write_json(out/'执行记录/启动日志.json',envelope(period,status='needs_confirmation',needs_reading=needs_reading,source_file_count=len(manifests)))
    return standardized, pre


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--root',required=True)
    p.add_argument('--period')
    p.add_argument('--files',nargs='*')
    p.add_argument('--output')
    p.add_argument('--mapping',required=True)
    p.add_argument('--currency',default='CNY')
    p.add_argument('--amount-precision',type=int)
    a=p.parse_args()
    _, pre=intake(a.root,a.period,a.files,a.output,read_json(a.mapping),a.currency,a.amount_precision)
    print(f"{pre['period_id']}: 预处理完成，正式事件须由 Agent 核查")


if __name__=='__main__':
    main()

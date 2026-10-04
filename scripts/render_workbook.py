"""Basic values-only Excel. The renderer receives exactly two audited objects."""
from __future__ import annotations

import argparse
import json
import math
import unicodedata
from datetime import datetime
from pathlib import Path
from string import Formatter

from common import VERSION, digest, period_dir, relative_file, resolve, sha, write_json
from validate_report import load_report_inputs, require, validate_inputs


def visible_sections(data,copy_text):
    for section in copy_text['sections']:
        condition=section['visible_when']
        if condition and resolve(data,condition['data_ref']) not in condition['status_in']:
            continue
        text=''
        for literal,field,fmt,_ in Formatter().parse(section['text']):
            text+=literal
            if field is not None:
                value=resolve(data,field)
                text+=copy_text['strings']['missing'] if value is None else format(value,fmt)
        yield text


def render_workbook(report_data, report_copy):
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    errors=validate_inputs(report_data,report_copy)
    require(not errors,'\n'.join(errors))
    data,copy_text=report_data,report_copy
    s=copy_text['strings']
    book=Workbook();book.remove(book.active)
    def table(name,headers,rows):
        sheet=book.create_sheet(s['sheet.'+name])
        sheet.sheet_view.showGridLines=False
        sheet.append([s['title']]);sheet.append([data['period_id'],data['currency']]);sheet.append([])
        sheet.append(headers)
        for row in rows:
            sheet.append(row)
        sheet['A1'].font=Font(name='Arial',size=16,bold=True,color='202838')
        for row in sheet.iter_rows(min_row=4):
            for cell in row:
                cell.font=Font(name='Arial',size=11,color='202838')
                cell.alignment=Alignment(vertical='top',wrap_text=True,horizontal='right' if isinstance(cell.value,(int,float)) else 'left')
                if isinstance(cell.value,str):
                    cell.data_type='s'  # CSV/Agent text such as =... stays text.
                if isinstance(cell.value,float):
                    cell.number_format='#,##0'+('.'+'0'*data['amount_precision'] if data['amount_precision'] else '')
            sheet.row_dimensions[row[0].row].height=42 if name in ('raw_record','normalized_record','economic_event','match_link','long_term') else 26
        for cell in sheet[4]:
            cell.fill=PatternFill('solid',fgColor='35465B');cell.font=Font(name='Arial',size=11,bold=True,color='FFFFFF')
        for col in range(1,len(headers)+1):
            width=32 if col==1 else 24
            if name in data['audit_tables']:
                key=data['audit_tables'][name]['columns'][col-1]
                if key in ('source_location','source_locations','record_allocations','raw_record_ids','record_ids'):
                    width=52
            sheet.column_dimensions[get_column_letter(col)].width=width
        if name in data['audit_tables']:
            for row in sheet.iter_rows(min_row=5):
                lines=1
                for cell in row:
                    width=sheet.column_dimensions[cell.column_letter].width
                    text=str(cell.value or '')
                    length=sum(2 if unicodedata.east_asian_width(ch) in ('W','F') else 1 for ch in text)
                    lines=max(lines,math.ceil(length/(width-2)))
                sheet.row_dimensions[row[0].row].height=max(42,lines*16+8)
        if rows:
            sheet.auto_filter.ref=f'A4:{get_column_letter(len(headers))}{sheet.max_row}'
        if len(rows)>12 or name in data['audit_tables']:
            sheet.freeze_panes='B5'
        return sheet
    rows=[]
    for key,m in data['metrics'].items():
        rows.append([s['metric.'+key],m['value'],s['status.'+m['status']],s['unit.'+m['unit']],m['currency']])
    overview=table('overview',[s[k] for k in ('header.metric','header.value','header.status','header.unit','header.currency')],rows)
    for i,(_,m) in enumerate(data['metrics'].items(),5):
        if m['unit']=='ratio':
            overview.cell(i,2).number_format='0.0%'
        if m['unit']=='currency_per_day':
            overview.cell(i,2).number_format='#,##0.0000'
    overview.append([])
    for key in ('expense','income','obligations'):
        overview.append([s['completeness.'+key],s['status.'+data['completeness'][key]]])
    for text in visible_sections(data,copy_text):
        overview.append([text])
        overview.row_dimensions[overview.max_row].height=48
    overview.column_dimensions['A'].width=44
    categories=[[c['category_l1'],c['amount'],c['count'],c['currency']] for c in data['categories']]
    table('categories',[s[k] for k in ('header.category','header.amount','header.count','header.currency')],categories)
    def cell_value(key,value):
        if isinstance(value,(dict,list)):
            return json.dumps(value,ensure_ascii=False,separators=(',',':'))
        if key=='trade_time' and value:
            return datetime.fromisoformat(value).replace(tzinfo=None)
        return value
    for name,t in data['audit_tables'].items():
        sheet=table(name,[s['field.'+k] for k in t['columns']],[[cell_value(k,r[k]) for k in t['columns']] for r in t['rows']])
        if 'trade_time' in t['columns']:
            col=t['columns'].index('trade_time')+1
            for row in range(5,sheet.max_row+1):
                sheet.cell(row,col).number_format='yyyy-mm-dd hh:mm:ss'
    long_rows=[[s['history.status'],s['status.'+data['comparison']['status']]],
               [s['history.samples'],data['comparison']['sample_count']],
               [s['history.periods'],', '.join(data['comparison']['reference_period_ids'])],
               [s['history.reference'],data['comparison']['reference_value']],
               [s['history.delta'],data['comparison']['delta']]]
    for g in data['goals']:
        long_rows.append([s['goal']+' '+g['goal_id'],s['goal_status.'+g['goal_status']],g['target_period'],g['actual_value']])
    for o in data['obligations']:
        long_rows.append([s['obligation']+' '+o['obligation_id'],s['obligation_status.'+o['obligation_status']],o['currency'],o['remaining_principal']])
    table('long_term',[s['header.item'],s['header.value'],s['header.period'],s['header.amount']],long_rows)
    for sheet in book:
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type='s'
    return book


def verify_workbook(path,data,copy_text):
    from openpyxl import load_workbook
    actual=load_workbook(path,data_only=False)
    expected=render_workbook(data,copy_text)
    try:
        require(actual.sheetnames==expected.sheetnames,'工作簿页结构不一致')
        for a,e in zip(actual.worksheets,expected.worksheets):
            av=list(a.values);ev=list(e.values)
            # Serialisation pads rows to the maximum used column.
            # XLSX stores numeric literals with 16 significant digits; compare that
            # documented export representation rather than binary float identity.
            ev=[tuple(None if v=='' else float(format(v,'.16g')) if isinstance(v,float) else v for v in row) for row in ev]
            require(av==ev,'工作簿读回数据不一致: '+a.title)
            require(a.auto_filter.ref==e.auto_filter.ref and a.freeze_panes==e.freeze_panes,'筛选/冻结设置丢失')
            require(not any(cell.data_type=='f' for row in a for cell in row),'基础工作簿不得依赖公式计算')
    finally:
        actual.close();expected.close()


def deliver(root,period):
    data,copy_text=load_report_inputs(root,period)
    year,month=period.split('.')
    base=period_dir(root,period)
    path=base/f'{year}年{month}月账单分析.xlsx'
    book=render_workbook(data,copy_text)
    book.save(path);book.close()
    verify_workbook(path,data,copy_text)
    audit=dict(schema_version=VERSION,period_id=period,report_data_sha256=digest(data),report_copy_sha256=digest(copy_text),
               workbook=path.relative_to(Path(root).resolve()).as_posix(),workbook_sha256=sha(path),status='passed')
    write_json(base/'执行记录/交付审计.json',audit)
    return path


def main():
    p=argparse.ArgumentParser(description=__doc__)
    p.add_argument('--root',required=True);p.add_argument('--period',required=True)
    a=p.parse_args();print(deliver(a.root,a.period))


if __name__=='__main__':
    main()
